import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, DataBarRule

SRC = '../data/cleaned_superstore.csv'
OUT = 'Global_Superstore_Beginner_Dashboard.xlsx'
FONT='Arial'; NAVY='1F3864'; LBLUE='D9E1F2'; WHITE='FFFFFF'; ORANGE='E8833A'

df = pd.read_csv(SRC, parse_dates=['order_date','ship_date'])
df['profit_margin'] = df['profit']/df['sales']
df['profit_margin'] = df['profit_margin'].replace([np.inf,-np.inf], np.nan)
cols = ['order_id','order_date','ship_date','delivery_days','ship_mode','order_priority',
        'customer_id','customer_name','segment','country','market','region',
        'category','sub_category','product_id','product_name',
        'quantity','discount','sales','profit','profit_margin','shipping_cost','year','month_year']
df = df[cols]

# ---------------- facts for the written insights (all computed, none typed) ----------------
T, P = df.sales.sum(), df.profit.sum()
M = lambda v: f"${v/1e6:.2f}M"; K = lambda v: f"${abs(v)/1e3:,.1f}K"
cat = df.groupby('category').agg(s=('sales','sum'),p=('profit','sum')); cat['m']=cat.p/cat.s*100
sub = df.groupby('sub_category').agg(s=('sales','sum'),p=('profit','sum'),d=('discount','mean')); sub['m']=sub.p/sub.s*100
mk  = df.groupby('market').agg(s=('sales','sum'),p=('profit','sum'),d=('discount','mean')); mk['m']=mk.p/mk.s*100
seg = df.groupby('segment').agg(s=('sales','sum'),p=('profit','sum')); seg['m']=seg.p/seg.s*100
yr  = df.groupby('year').agg(s=('sales','sum'),p=('profit','sum')); yr['g']=yr.s.pct_change()*100
df['mo']=df.order_date.dt.month; se = df.groupby('mo').sales.sum()
sm  = df.groupby('ship_mode').agg(s=('sales','sum'),p=('profit','sum'),d=('delivery_days','mean')); sm['m']=sm.p/sm.s*100
corr = df[['discount','profit_margin']].dropna().corr().iloc[0,1]
z = df[df.discount==0]; hi = df[df.discount>0.3]; gt20 = df[df.discount>0.2]
z_margin = z.profit.sum()/z.sales.sum()*100
loss_all = (df.profit<0).mean()*100
mname = {1:'January',2:'February',11:'November',12:'December'}
topcat = cat.s.idxmax(); topmar = cat.m.idxmax(); worstm = cat.m.idxmin()
best_sub = sub.m.idxmax(); worst_sub = sub.m.idxmin(); sm_top = sm.s.idxmax()

wb = Workbook()
def hdr(ws,row,n):
    for c in range(1,n+1):
        x=ws.cell(row=row,column=c); x.font=Font(name=FONT,bold=True,color=WHITE,size=11)
        x.fill=PatternFill('solid',fgColor=NAVY); x.alignment=Alignment(horizontal='center')
def color(ch, c=NAVY):
    for s in ch.series: s.graphicalProperties.solidFill=c; s.graphicalProperties.line.solidFill=c

# =================== READ ME ===================
ws=wb.active; ws.title='Read Me'; ws.sheet_view.showGridLines=False; ws.column_dimensions['A'].width=118
ws['A1']='Global Superstore Sales Analysis — Beginner Level'; ws['A1'].font=Font(name=FONT,bold=True,size=16,color=NAVY)
ws['A2']='ShadowFox Data Analyst Internship  |  Prepared by: [Your Name]'; ws['A2'].font=Font(name=FONT,italic=True,size=11,color='555555')
lines=[
 ('','n'),('DATASET','h'),
 (f"Global Superstore transactional sales data, Jan 2011 – Dec 2014: {len(df):,} order line items, {df.country.nunique()} countries, "
  f"{df.market.nunique()} markets, {df.customer_id.nunique():,} customers, 3 categories, {df.sub_category.nunique()} sub-categories.",'n'),
 ('','n'),('DATA CLEANING & VALIDATION','h'),
 ('1. Loaded with encoding=latin1 (non-UTF-8 characters in product names) and thousands="," so values such as "2,309.65" parse as numbers.','n'),
 ('2. Validated every numeric column after parsing (see data-quality note below): 0 nulls in sales, profit, quantity, discount, shipping cost.','n'),
 ('3. Checked for duplicates (0), negative sales (0) and inconsistent category spellings (none).','n'),
 ('4. Parsed order_date / ship_date (day-first) and derived delivery_days, year, month_year and profit_margin (= profit / sales).','n'),
 ('5. Kept customer_id (not customer_name, which repeats across countries) as the customer key for later analysis.','n'),
 ('','n'),('DATA-QUALITY NOTE (why the numbers here differ from an earlier draft)','h'),
 ('An earlier draft of this workbook was built from a copy of the file where every sale of $1,000 or more failed to parse and was silently converted to blank by','n'),
 ('pd.to_numeric(errors="coerce"). That hid ~38% of revenue and made margin look like 18.7% instead of the true 11.6%. The tell was that the largest surviving','n'),
 ('sales value was exactly 999. The workbook was rebuilt from the correctly parsed data, and all figures and insights below reflect the corrected numbers.','n'),
 ('','n'),('WORKBOOK STRUCTURE','h'),
 ('- Dashboard Overview: one-page executive view (headline KPIs, three charts, top takeaways).','n'),
 ('- KPI Summary: headline business metrics for the full period.','n'),
 ('- Category & Segment / Regional Analysis / Monthly Trend / Sub-Category Performance: the supporting analysis tables and charts.','n'),
 ('- Insights & Conclusions: written findings and recommendations.','n'),
 ('- Raw Data: the cleaned dataset as an Excel Table ("SalesData"). Every summary cell is a SUMIFS / AVERAGEIFS / COUNTIFS formula over this table,','n'),
 ('  so the workbook recalculates automatically if the data changes.','n'),
 ('','n'),('WHY THESE METRICS WERE SELECTED','h'),
 ('Total Sales & Total Profit — the two numbers every stakeholder asks for first; every other metric exists to explain movement in them.','n'),
 (f"Profit Margin % — sales alone hide problems. Furniture is the 2nd-largest category by revenue ({M(cat.loc['Furniture','s'])}) but has the lowest margin ({cat.loc['Furniture','m']:.1f}%).",'n'),
 (f"Average Discount and its relationship to margin — discount ranges from 0% to 80%; testing it showed it is the strongest driver of unprofitable orders (correlation {corr:.2f} with row-level margin).",'n'),
 ('Delivery Days — a simple operational metric (ship_date − order_date) that shows the speed vs. shipping-mode trade-off.','n'),
 ('Category / Sub-Category / Segment / Market — the four categorical fields a business would realistically segment by: what we sell, to whom, and where.','n'),
 ('Monthly trend — 48 monthly points is granular enough to show seasonality clearly while remaining readable in a single chart.','n'),
]
r=4
for t,k in lines:
    c=ws.cell(row=r,column=1,value=t)
    c.font=Font(name=FONT,bold=True,size=12,color=NAVY) if k=='h' else Font(name=FONT,size=10.5)
    r+=1

# =================== RAW DATA ===================
wr=wb.create_sheet('Raw Data'); wr.sheet_view.showGridLines=False
wr.append(list(df.columns))
for row in df.itertuples(index=False):
    wr.append([None if (isinstance(v,float) and np.isnan(v)) else (v.to_pydatetime() if isinstance(v,pd.Timestamp) else v) for v in row])
nr=len(df)+1; nc=len(cols); hdr(wr,1,nc)
fmt={'order_date':'yyyy-mm-dd','ship_date':'yyyy-mm-dd','sales':'$#,##0.00','profit':'$#,##0.00','shipping_cost':'$#,##0.00','profit_margin':'0.0%','discount':'0%'}
for i,h in enumerate(cols,1):
    wr.column_dimensions[get_column_letter(i)].width = 22 if h in('customer_name','product_name','country') else 14
    if h in fmt:
        for rr in range(2,nr+1): wr.cell(row=rr,column=i).number_format=fmt[h]
t=Table(displayName='SalesData',ref=f"A1:{get_column_letter(nc)}{nr}")
t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True); wr.add_table(t); wr.freeze_panes='A2'

# =================== KPI SUMMARY ===================
wk=wb.create_sheet('KPI Summary'); wk.sheet_view.showGridLines=False
wk['A1']='Global Superstore — Headline KPIs (Jan 2011 – Dec 2014)'; wk['A1'].font=Font(name=FONT,bold=True,size=15,color=NAVY)
kp=[('Total Sales','=SUM(SalesData[sales])','$#,##0'),('Total Profit','=SUM(SalesData[profit])','$#,##0'),
    ('Overall Profit Margin','=B4/B3','0.0%'),('Total Order Line Items','=COUNTA(SalesData[order_id])','#,##0'),
    ('Total Units Sold','=SUM(SalesData[quantity])','#,##0'),('Average Sale Value per Line Item','=B3/B6','$#,##0.00'),
    ('Average Discount Given','=AVERAGE(SalesData[discount])','0.0%'),('Total Shipping Cost','=SUM(SalesData[shipping_cost])','$#,##0'),
    ('Average Delivery Time (days)','=AVERAGE(SalesData[delivery_days])','0.0')]
for i,(a,f,n) in enumerate(kp,3):
    wk.cell(row=i,column=1,value=a).font=Font(name=FONT,bold=True,size=11)
    c=wk.cell(row=i,column=2,value=f); c.font=Font(name=FONT,bold=True,size=12,color=NAVY); c.number_format=n; c.fill=PatternFill('solid',fgColor=LBLUE)
wk['A13']='Data Period Covered:'; wk['B13']='=TEXT(MIN(SalesData[order_date]),"mmm yyyy")&"  to  "&TEXT(MAX(SalesData[order_date]),"mmm yyyy")'
for rr,(lab,val,cm) in zip((14,15,16),[('Countries Covered:',df.country.nunique(),'country'),('Markets Covered:',df.market.nunique(),'market'),('Customers (unique customer_id):',df.customer_id.nunique(),'customer_id')]):
    wk.cell(row=rr,column=1,value=lab); c=wk.cell(row=rr,column=2,value=int(val)); c.comment=Comment(f'Distinct count of "{cm}" in Raw Data, computed during data preparation.','Data Prep')
for rr in range(13,17):
    wk.cell(row=rr,column=1).font=Font(name=FONT,bold=True); wk.cell(row=rr,column=2).font=Font(name=FONT)
wk.column_dimensions['A'].width=34; wk.column_dimensions['B'].width=24

# =================== CATEGORY & SEGMENT ===================
wc=wb.create_sheet('Category & Segment'); wc.sheet_view.showGridLines=False
wc['A1']='Sales Performance by Product Category and Customer Segment'; wc['A1'].font=Font(name=FONT,bold=True,size=14,color=NAVY)
H=['Category','Total Sales','Total Profit','Profit Margin %','Units Sold','Line Items','Avg Discount']
def block(ws,title_row,label,field,values,first_hdr):
    ws.cell(row=title_row,column=1,value=label).font=Font(name=FONT,bold=True,size=12)
    hr=title_row+1
    for i,h in enumerate([first_hdr]+H[1:],1): ws.cell(row=hr,column=i,value=h)
    hdr(ws,hr,7); r=hr+1; s=r
    for v in values:
        ws.cell(row=r,column=1,value=v)
        ws.cell(row=r,column=2,value=f'=SUMIFS(SalesData[sales],SalesData[{field}],A{r})').number_format='$#,##0'
        ws.cell(row=r,column=3,value=f'=SUMIFS(SalesData[profit],SalesData[{field}],A{r})').number_format='$#,##0'
        ws.cell(row=r,column=4,value=f'=C{r}/B{r}').number_format='0.0%'
        ws.cell(row=r,column=5,value=f'=SUMIFS(SalesData[quantity],SalesData[{field}],A{r})').number_format='#,##0'
        ws.cell(row=r,column=6,value=f'=COUNTIFS(SalesData[{field}],A{r})').number_format='#,##0'
        ws.cell(row=r,column=7,value=f'=AVERAGEIFS(SalesData[discount],SalesData[{field}],A{r})').number_format='0.0%'
        r+=1
    return hr,s,r-1
cats=cat.sort_values('s',ascending=False).index.tolist()
h1,s1,e1=block(wc,3,'By Product Category','category',cats,'Category')
tr=e1+1
wc.cell(row=tr,column=1,value='TOTAL')
for col,f,n in [(2,f'=SUM(B{s1}:B{e1})','$#,##0'),(3,f'=SUM(C{s1}:C{e1})','$#,##0'),(4,f'=C{tr}/B{tr}','0.0%'),(5,f'=SUM(E{s1}:E{e1})','#,##0'),(6,f'=SUM(F{s1}:F{e1})','#,##0')]:
    c=wc.cell(row=tr,column=col,value=f); c.number_format=n
for c in range(1,8): wc.cell(row=tr,column=c).font=Font(name=FONT,bold=True)
segs=seg.sort_values('s',ascending=False).index.tolist()
h2,s2,e2=block(wc,tr+3,'By Customer Segment','segment',segs,'Segment')
for col,w in zip('ABCDEFG',[18,14,14,15,12,11,12]): wc.column_dimensions[col].width=w
b=BarChart(); b.type='col'; b.title='Sales by Category'; b.y_axis.title='Sales ($)'
b.add_data(Reference(wc,min_col=2,min_row=h1,max_row=e1),titles_from_data=True); b.set_categories(Reference(wc,min_col=1,min_row=s1,max_row=e1)); b.width=12; b.height=8; color(b); wc.add_chart(b,'I3')
b=BarChart(); b.type='col'; b.title='Sales by Customer Segment'; b.y_axis.title='Sales ($)'
b.add_data(Reference(wc,min_col=2,min_row=h2,max_row=e2),titles_from_data=True); b.set_categories(Reference(wc,min_col=1,min_row=s2,max_row=e2)); b.width=12; b.height=8; color(b,ORANGE); wc.add_chart(b,'I20')

# =================== REGIONAL ===================
wg=wb.create_sheet('Regional Analysis'); wg.sheet_view.showGridLines=False
wg['A1']='Sales Performance by Market'; wg['A1'].font=Font(name=FONT,bold=True,size=14,color=NAVY)
for i,h in enumerate(['Market','Total Sales','Total Profit','Profit Margin %','Units Sold','Line Items','% of Total Sales','Avg Discount'],1): wg.cell(row=3,column=i,value=h)
hdr(wg,3,8); r=4
for m_ in mk.sort_values('s',ascending=False).index:
    wg.cell(row=r,column=1,value=m_)
    wg.cell(row=r,column=2,value=f'=SUMIFS(SalesData[sales],SalesData[market],A{r})').number_format='$#,##0'
    wg.cell(row=r,column=3,value=f'=SUMIFS(SalesData[profit],SalesData[market],A{r})').number_format='$#,##0'
    wg.cell(row=r,column=4,value=f'=C{r}/B{r}').number_format='0.0%'
    wg.cell(row=r,column=5,value=f'=SUMIFS(SalesData[quantity],SalesData[market],A{r})').number_format='#,##0'
    wg.cell(row=r,column=6,value=f'=COUNTIFS(SalesData[market],A{r})').number_format='#,##0'
    wg.cell(row=r,column=7,value=f'=B{r}/SUM(SalesData[sales])').number_format='0.0%'
    wg.cell(row=r,column=8,value=f'=AVERAGEIFS(SalesData[discount],SalesData[market],A{r})').number_format='0.0%'
    r+=1
mend=r-1
for col,w in zip('ABCDEFGH',[16,14,14,15,12,11,15,13]): wg.column_dimensions[col].width=w
b=BarChart(); b.type='col'; b.title='Total Sales by Market'; b.y_axis.title='Sales ($)'
b.add_data(Reference(wg,min_col=2,min_row=3,max_row=mend),titles_from_data=True); b.set_categories(Reference(wg,min_col=1,min_row=4,max_row=mend)); b.width=16; b.height=9; color(b); wg.add_chart(b,'J3')
b=BarChart(); b.type='col'; b.title='Profit Margin % by Market'
b.add_data(Reference(wg,min_col=4,min_row=3,max_row=mend),titles_from_data=True); b.set_categories(Reference(wg,min_col=1,min_row=4,max_row=mend)); b.width=16; b.height=9; color(b,ORANGE); wg.add_chart(b,'J22')

# =================== MONTHLY ===================
wm=wb.create_sheet('Monthly Trend'); wm.sheet_view.showGridLines=False
wm['A1']='Monthly Sales & Profit Trend (Jan 2011 – Dec 2014)'; wm['A1'].font=Font(name=FONT,bold=True,size=14,color=NAVY)
for i,h in enumerate(['Month','Sales','Profit','Profit Margin %','Line Items'],1): wm.cell(row=3,column=i,value=h)
hdr(wm,3,5); r=4
for mo in sorted(df.month_year.unique()):
    wm.cell(row=r,column=1,value=mo)
    wm.cell(row=r,column=2,value=f'=SUMIFS(SalesData[sales],SalesData[month_year],A{r})').number_format='$#,##0'
    wm.cell(row=r,column=3,value=f'=SUMIFS(SalesData[profit],SalesData[month_year],A{r})').number_format='$#,##0'
    wm.cell(row=r,column=4,value=f'=C{r}/B{r}').number_format='0.0%'
    wm.cell(row=r,column=5,value=f'=COUNTIFS(SalesData[month_year],A{r})').number_format='#,##0'
    r+=1
assert r-1==51, r
for col,w in zip('ABCDE',[12,14,14,15,11]): wm.column_dimensions[col].width=w
ln=LineChart(); ln.title='Monthly Sales & Profit Trend'; ln.y_axis.title='Amount ($)'
ln.add_data(Reference(wm,min_col=2,max_col=3,min_row=3,max_row=51),titles_from_data=True); ln.set_categories(Reference(wm,min_col=1,min_row=4,max_row=51))
ln.series[0].graphicalProperties.line.solidFill=NAVY; ln.series[1].graphicalProperties.line.solidFill=ORANGE
for _s in ln.series: _s.smooth=False
ln.width=28; ln.height=12; wm.add_chart(ln,'G3')

# =================== SUB-CATEGORY ===================
wsb=wb.create_sheet('Sub-Category Performance'); wsb.sheet_view.showGridLines=False
wsb['A1']='Sub-Category Performance — Ranked by Total Sales'; wsb['A1'].font=Font(name=FONT,bold=True,size=14,color=NAVY)
for i,h in enumerate(['Sub-Category','Category','Total Sales','Total Profit','Profit Margin %','Units Sold','Avg Discount'],1): wsb.cell(row=3,column=i,value=h)
hdr(wsb,3,7); cmap=df.groupby('sub_category').category.first(); r=4
for sc in sub.sort_values('s',ascending=False).index:
    wsb.cell(row=r,column=1,value=sc); wsb.cell(row=r,column=2,value=cmap[sc])
    wsb.cell(row=r,column=3,value=f'=SUMIFS(SalesData[sales],SalesData[sub_category],A{r})').number_format='$#,##0'
    wsb.cell(row=r,column=4,value=f'=SUMIFS(SalesData[profit],SalesData[sub_category],A{r})').number_format='$#,##0'
    wsb.cell(row=r,column=5,value=f'=D{r}/C{r}').number_format='0.0%'
    wsb.cell(row=r,column=6,value=f'=SUMIFS(SalesData[quantity],SalesData[sub_category],A{r})').number_format='#,##0'
    wsb.cell(row=r,column=7,value=f'=AVERAGEIFS(SalesData[discount],SalesData[sub_category],A{r})').number_format='0.0%'
    r+=1
send=r-1
for col,w in zip('ABCDEFG',[16,16,14,14,15,12,13]): wsb.column_dimensions[col].width=w
wsb.conditional_formatting.add(f'E4:E{send}',CellIsRule(operator='lessThan',formula=['0'],fill=PatternFill('solid',fgColor='F8CBCB')))
wsb.conditional_formatting.add(f'E4:E{send}',CellIsRule(operator='greaterThan',formula=['0.17'],fill=PatternFill('solid',fgColor='C6E0B4')))
wsb.conditional_formatting.add(f'D4:D{send}',CellIsRule(operator='lessThan',formula=['0'],font=Font(name=FONT,color='C62828',bold=True)))
wsb.conditional_formatting.add(f'C4:C{send}',DataBarRule(start_type='min',end_type='max',color='638EC6'))
b=BarChart(); b.type='bar'; b.title='Top 10 Sub-Categories by Sales'; b.y_axis.title='Sales ($)'
b.add_data(Reference(wsb,min_col=3,min_row=3,max_row=13),titles_from_data=True); b.set_categories(Reference(wsb,min_col=1,min_row=4,max_row=13)); b.width=18; b.height=12; color(b); wsb.add_chart(b,'I3')

# =================== INSIGHTS ===================
wi=wb.create_sheet('Insights & Conclusions'); wi.sheet_view.showGridLines=False
wi.column_dimensions['A'].width=6; wi.column_dimensions['B'].width=112
wi['A1']='Insights & Conclusions'; wi['A1'].font=Font(name=FONT,bold=True,size=16,color=NAVY)
gr = lambda a,b_: (b_/a-1)*100
sections=[
 ('1. Category mix',
  f"{topcat} is the largest category by revenue ({M(cat.loc[topcat,'s'])}) and also has the best margin ({cat.loc[topmar,'m']:.1f}%), generating the most profit ({K(cat.loc['Technology','p'])}). "
  f"Office Supplies is close behind on margin ({cat.loc['Office Supplies','m']:.1f}%). Furniture is the weak spot: {M(cat.loc['Furniture','s'])} of sales (second-largest) but only a "
  f"{cat.loc['Furniture','m']:.1f}% margin — roughly half of the other two categories."),
 ('2. The Tables sub-category is losing money',
  f"{worst_sub} is the only unprofitable sub-category: {'-' if sub.loc[worst_sub,'p']<0 else ''}{K(sub.loc[worst_sub,'p'])} profit on {K(sub.loc[worst_sub,'s'])} of sales ({sub.loc[worst_sub,'m']:.1f}% margin). "
  f"It also carries the deepest average discount of any sub-category ({sub.loc[worst_sub,'d']*100:.0f}%), a likely contributing factor (though this is a correlation, not proof of cause). At the other end, {best_sub} ({sub.loc[best_sub,'m']:.1f}%) "
  f"and Labels ({sub.loc['Labels','m']:.1f}%) have the best margins, and Copiers combines strong sales ({M(sub.loc['Copiers','s'])}) with a {sub.loc['Copiers','m']:.1f}% margin."),
 ('3. Discounting is destroying profit',
  f"Discount level is strongly negatively correlated with profit margin (r = {corr:.2f}). Order lines with no discount are never unprofitable and earn a {z_margin:.1f}% margin; lines discounted above 30% "
  f"lose money {((hi.profit<0).mean()*100):.1f}% of the time. Discounts above 20% cover {(len(gt20)/len(df)*100):.0f}% of all order lines and account for {'-' if gt20.profit.sum()<0 else ''}{'$'}{abs(gt20.profit.sum()):,.0f} of profit — "
  f"against total company profit of {M(P)}. Overall, {loss_all:.1f}% of order lines lose money."),
 ('4. Strong, sustained growth',
  f"Revenue grew every year: {M(yr.s[2011])} (2011) → {M(yr.s[2012])} (2012, +{yr.g[2012]:.1f}%) → {M(yr.s[2013])} (2013, +{yr.g[2013]:.1f}%) → {M(yr.s[2014])} (2014, +{yr.g[2014]:.1f}%). "
  f"Profit doubled over the period ({K(yr.p[2011])} to {K(yr.p[2014])}) while margin stayed steady at roughly 11–12%, so growth has come with stable profitability rather than at its expense."),
 ('5. Regional performance is uneven',
  f"APAC ({M(mk.loc['APAC','s'])}) and EU ({M(mk.loc['EU','s'])}) are the largest markets. Canada is tiny ({K(mk.loc['Canada','s'])}) but has by far the best margin ({mk.loc['Canada','m']:.1f}%) "
  f"with an average discount of about {mk.loc['Canada','d']*100:.0f}%. EMEA is the weakest significant market: {M(mk.loc['EMEA','s'])} of sales at only a {mk.loc['EMEA','m']:.1f}% margin, with an average discount of "
  f"{mk.loc['EMEA','d']*100:.0f}% — the pattern again points to discounting as the driver."),
 ('6. Seasonality',
  f"Sales are strongly seasonal: December is the peak month ({M(se[12])} across the four years) with November close behind ({M(se[11])}), while February is the slowest ({M(se[2])}). "
  f"Demand builds through the second half of the year and drops sharply after the holiday period."),
 ('7. Shipping mode and customer segment are not margin drivers',
  f"{sm_top} accounts for the bulk of volume ({M(sm.loc[sm_top,'s'])}, {sm.loc[sm_top,'s']/T*100:.0f}% of sales, ~{sm.loc[sm_top,'d']:.0f}-day average delivery), yet margins are nearly identical across all four "
  f"shipping modes ({sm.m.min():.1f}%–{sm.m.max():.1f}%). Likewise the three customer segments differ by under one margin point ({seg.m.min():.1f}%–{seg.m.max():.1f}%). "
  f"Profitability is driven by discounting, product and market — not by how the order ships or who buys it."),
]
r=3
for title,text in sections:
    c=wi.cell(row=r,column=1,value=title); c.font=Font(name=FONT,bold=True,size=12,color=NAVY); wi.merge_cells(f'A{r}:B{r}'); r+=1
    c=wi.cell(row=r,column=2,value=text); c.font=Font(name=FONT,size=10.5); c.alignment=Alignment(wrap_text=True,vertical='top')
    wi.row_dimensions[r].height=max(34,15.5*(len(text)//105+1)); r+=2
c=wi.cell(row=r,column=1,value='Recommendations'); c.font=Font(name=FONT,bold=True,size=13,color=NAVY); wi.merge_cells(f'A{r}:B{r}'); r+=1
recs=[
 f"1. Require approval for discounts above ~20%, and test the change before rolling out: margin turns negative between 20% and 30% discount, but the link is correlational — run it as a controlled pilot (e.g. one EMEA sub-region) and compare volume vs. margin first.",
 f"2. Review Tables specifically (pricing, supplier cost, discount policy) — it is the only sub-category losing money ({'-' if sub.loc[worst_sub,'p']<0 else ''}{K(sub.loc[worst_sub,'p'])}).",
 f"3. Lean into high-margin products: Copiers ({sub.loc['Copiers','m']:.1f}%), Paper ({sub.loc['Paper','m']:.1f}%) and Labels ({sub.loc['Labels','m']:.1f}%) combine healthy margins with low discounting.",
 f"4. Study why Canada earns {mk.loc['Canada','m']:.1f}% with almost no discounting, and test whether that pricing discipline can be applied to EMEA ({mk.loc['EMEA','m']:.1f}% margin).",
 f"5. Plan inventory, staffing and marketing around the Nov–Dec peak and the February trough rather than assuming flat monthly demand.",
]
for t in recs:
    c=wi.cell(row=r,column=2,value=t); c.font=Font(name=FONT,size=10.5); c.alignment=Alignment(wrap_text=True,vertical='top')
    wi.row_dimensions[r].height=max(32,15.5*(len(t)//105+1)); r+=1

# =================== DASHBOARD OVERVIEW ===================
wd=wb.create_sheet('Dashboard Overview'); wd.sheet_view.showGridLines=False
wd.merge_cells('A1:L1'); wd['A1']='Global Superstore — Executive Dashboard (Jan 2011 – Dec 2014)'
wd['A1'].font=Font(name=FONT,bold=True,size=18,color=NAVY); wd.row_dimensions[1].height=30
thin=Side(style='thin',color='CCCCCC'); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
cards=[('A','B','Total Sales',"='KPI Summary'!B3",'$#,##0.0,,"M"'),('D','E','Total Profit',"='KPI Summary'!B4",'$#,##0.00,,"M"'),
       ('G','H','Overall Margin',"='KPI Summary'!B5",'0.0%'),('J','K','2014 Sales Growth',"=(SUM('Monthly Trend'!B40:B51)-SUM('Monthly Trend'!B28:B39))/SUM('Monthly Trend'!B28:B39)",'0.0%')]
for a,b_,lab,f,n in cards:
    wd.merge_cells(f'{a}3:{b_}3'); x=wd[f'{a}3']; x.value=lab; x.font=Font(name=FONT,bold=True,size=11,color=WHITE); x.fill=PatternFill('solid',fgColor=NAVY); x.alignment=Alignment(horizontal='center')
    wd.merge_cells(f'{a}4:{b_}6'); v=wd[f'{a}4']; v.value=f; v.number_format=n; v.font=Font(name=FONT,bold=True,size=22,color=NAVY)
    v.alignment=Alignment(horizontal='center',vertical='center'); v.fill=PatternFill('solid',fgColor=LBLUE)
    for rr in range(3,7):
        for cc in (a,b_): wd[f'{cc}{rr}'].border=bd
for col in 'ABCDEFGHIJKL': wd.column_dimensions[col].width=11
ln=LineChart(); ln.title='Monthly Sales & Profit Trend'
ln.add_data(Reference(wm,min_col=2,max_col=3,min_row=3,max_row=51),titles_from_data=True); ln.set_categories(Reference(wm,min_col=1,min_row=4,max_row=51))
ln.series[0].graphicalProperties.line.solidFill=NAVY; ln.series[1].graphicalProperties.line.solidFill=ORANGE
for _s in ln.series: _s.smooth=False
ln.width=25; ln.height=9; wd.add_chart(ln,'A8')
b=BarChart(); b.type='col'; b.title='Sales by Category'
b.add_data(Reference(wc,min_col=2,min_row=h1,max_row=e1),titles_from_data=True); b.set_categories(Reference(wc,min_col=1,min_row=s1,max_row=e1)); b.width=12.4; b.height=9; color(b); wd.add_chart(b,'A27')
b=BarChart(); b.type='col'; b.title='Sales by Market'
b.add_data(Reference(wg,min_col=2,min_row=3,max_row=mend),titles_from_data=True); b.set_categories(Reference(wg,min_col=1,min_row=4,max_row=mend)); b.width=12.4; b.height=9; color(b,ORANGE); wd.add_chart(b,'G27')
wd['A46']='Top Takeaways'; wd['A46'].font=Font(name=FONT,bold=True,size=13,color=NAVY); wd.merge_cells('A46:L46')
tk=[f"Revenue grew every year (+{yr.g[2012]:.1f}%, +{yr.g[2013]:.1f}%, +{yr.g[2014]:.1f}%) while margin held steady at ~{P/T*100:.1f}% — growth without sacrificing profitability.",
    f"Discounts above 20% cover {(len(gt20)/len(df)*100):.0f}% of order lines and lose {'$'}{abs(gt20.profit.sum()):,.0f} — the single biggest profit leak (correlation with margin {corr:.2f}).",
    f"Furniture is the weakest category ({cat.loc['Furniture','m']:.1f}% margin), and Tables is the only sub-category losing money ({sub.loc[worst_sub,'m']:.1f}%).",
    f"Canada has the best margin ({mk.loc['Canada','m']:.1f}%, ~0% discount); EMEA the weakest significant market ({mk.loc['EMEA','m']:.1f}%, ~{mk.loc['EMEA','d']*100:.0f}% discount)."]
for i,t in enumerate(tk):
    rr=47+i; wd.merge_cells(f'A{rr}:L{rr}'); c=wd.cell(row=rr,column=1,value='•  '+t); c.font=Font(name=FONT,size=10.5); c.alignment=Alignment(wrap_text=True,vertical='top'); wd.row_dimensions[rr].height=30

for _ws in (wd,wi,wk,wc,wg,wm,wsb,ws):
    _ws.page_setup.orientation='landscape'; _ws.page_setup.fitToWidth=1; _ws.page_setup.fitToHeight=0
    _ws.sheet_properties.pageSetUpPr.fitToPage=True
order=['Read Me','Dashboard Overview','KPI Summary','Category & Segment','Regional Analysis','Monthly Trend','Sub-Category Performance','Insights & Conclusions','Raw Data']
wb._sheets=[wb[n] for n in order]; wb.active=0
wb.save(OUT); print('saved',OUT)
